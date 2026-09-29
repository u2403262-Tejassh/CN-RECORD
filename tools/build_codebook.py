#!/usr/bin/env python3
"""Build the lab codebook from Java sources, PNG outputs and Familiarisation.docx.
Run: .venv/bin/python tools/build_codebook.py
The only new experiment source is Day 1/network_setup.sh. Familiarization
pages 9–11 contain study material and results, without screenshot placeholders.
"""
from __future__ import annotations

from io import BytesIO
from math import ceil
from pathlib import Path
import subprocess
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
DEST = ROOT / 'Computer_Networks_Lab_Codebook.docx'
BLACK = RGBColor(22, 22, 22)
GRAY = BLACK  # The supplied record format specifies black text throughout.
DATES = ['18/07/2026', '18/07/2026', '15/07/2026', '05/08/2026',
         '19/08/2026', '02/09/2026', '09/09/2026', '09/09/2026',
         '01/07/2026', '08/07/2026', '01/07/2026']
TITLES = [
    'Familiarization: Linux networking and automatic network configuration',
    'Client–server communication using TCP socket programming',
    'Client–server communication using UDP socket programming',
    'ARQ flow control: Stop-and-Wait, Go-Back-N and Selective Repeat',
    'Implementation and simulation of Distance Vector Routing',
    'Implementation and simulation of the Leaky Bucket algorithm',
    'Implementation of Simple Mail Transfer Protocol (SMTP)',
    'Implementation of File Transfer Protocol (FTP)',
    'Introduction to the Wireshark tool',
    'Familiarization with the NS2 simulator',
    'Study of Packet Tracer',
]
AIMS = [
    'Inspect Linux network configuration and automate IPv4 address, default gateway, connectivity test and routing capture.',
    'Exchange data between a TCP server and client, then explore chat and remote-calculator variants.',
    'Exchange a datagram and a response using UDP sockets.',
    'Observe retransmission after frame loss using three ARQ simulations.',
    'Calculate routing costs and next hops from an input cost matrix.',
    'Simulate a finite-capacity bucket with a fixed output rate and overflow queued for resend.',
    'Exchange SMTP-style greetings, envelope, message data and termination over a local socket.',
    'Transfer a text file from a client to a server over a local TCP connection.',
    'Learn to capture, filter and inspect network traffic with Wireshark.',
    'Understand topology creation, Tcl scenarios, trace files and NAM in NS2.',
    'Learn to assemble, configure and test a small virtual network in Packet Tracer.',
]
# A maximum of 49 source lines per code page leaves comfortable room for the
# title/aim/running note and avoids dependence on a particular Word renderer.
MAX_LINES = 49


def parts(src: str):
    lines = (ROOT / src).read_text(encoding='utf-8').expandtabs(4).splitlines()
    count = ceil(len(lines) / MAX_LINES)
    per_page = ceil(len(lines) / count)
    return ['\n'.join(lines[i:i + per_page]) for i in range(0, len(lines), per_page)]


def filepages(day, paths):
    out = []
    for path in paths:
        chunks = parts(path)
        for i, text in enumerate(chunks):
            out.append({'day': day, 'kind': 'code', 'file': path, 'part': i + 1,
                        'total': len(chunks), 'text': text})
    return out


def outputs(day, groups):
    return [{'day': day, 'kind': 'images', 'items': group} for group in groups]


PAGES = [{'day': 1, 'kind': 'linux'}]
PAGES += filepages(1, ['Day 1/network_setup.sh'])
PAGES += [{'day': 1, 'kind': 'dryrun'}]
PAGES += filepages(2, ['TCP UDP Socket/TCPServer.java', 'TCP UDP Socket/TCPClient.java',
                       'TCP UDP Socket/TCPServerChat.java', 'TCP UDP Socket/TCPClientChat.java',
                       'TCP UDP Socket/TCPServerCalc.java', 'TCP UDP Socket/TCPClientCalc.java'])
PAGES += outputs(2, [
    [('TCP UDP Socket/TCPServer.png', 'TCP server'), ('TCP UDP Socket/TCPClient.png', 'TCP client')],
    [('TCP UDP Socket/TCPServerChat.png', 'Chat server'), ('TCP UDP Socket/TCPClientChat.png', 'Chat client')],
    [('TCP UDP Socket/TCPServerCalc.png', 'Calculator server'), ('TCP UDP Socket/TCPClientCalc.png', 'Calculator client')],
])
PAGES += filepages(3, ['TCP UDP Socket/UDPServer.java', 'TCP UDP Socket/UDPClient.java'])
PAGES += outputs(3, [[('TCP UDP Socket/UDPServer.png', 'UDP server'),
                      ('TCP UDP Socket/UDPClient.png', 'UDP client')]])
PAGES += filepages(4, ['ARQ/StopAndWait.java', 'ARQ/GoBackN.java', 'ARQ/SelectiveRepeat.java'])
PAGES += outputs(4, [[('ARQ/StopAndWait.png', 'Stop-and-Wait ARQ')],
                     [('ARQ/GoBackN.png', 'Go-Back-N ARQ')],
                     [('ARQ/SelectiveRepeat.png', 'Selective Repeat ARQ')]])
PAGES += filepages(5, ['DVR/DistanceVector.java'])
PAGES += outputs(5, [[('DVR/day6.png', 'DistanceVector — routing tables')]])
PAGES += filepages(6, ['ARQ/LeakyBucket.java', 'ARQ/LeakyBucketv2.java'])
PAGES += outputs(6, [[('ARQ/LeakyBucket.png', 'LeakyBucket')],
                     [('ARQ/LeakyBucketv2.png', 'LeakyBucketv2')]])
PAGES += filepages(7, ['SMTP/SMTPServer.java', 'SMTP/SMTPClient.java'])
PAGES += outputs(7, [[('SMTP/SMTPServer.png', 'SMTP server'),
                      ('SMTP/SMTPClient.png', 'SMTP client')]])
PAGES += filepages(8, ['FTP/FTPServer.java', 'FTP/FTPClient.java'])
PAGES += outputs(8, [[('FTP/FTP1.png', 'FTP server'), ('FTP/FTP2.png', 'FTP client')]])
PAGES += [{'day': n, 'kind': 'familiarization'} for n in (9, 10, 11)]
STARTS = {day: 1 + next(i for i, page in enumerate(PAGES) if page['day'] == day)
          for day in range(1, 12)}


def font(run, size=12, bold=False, color=BLACK, face='Times New Roman'):
    run.font.name = face
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    # Explicit East Asian font prevents arbitrary fallback on other machines.
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is not None:
        rfonts.set(qn('w:eastAsia'), face)
    return run


def para(doc, text='', size=12, bold=False, after=5, before=0, align=None,
         line=1.5, color=BLACK):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    font(p.add_run(text), size, bold, color)
    return p


def label(doc, text, before=8):
    return para(doc, text, 10, True, after=4, before=before, align=WD_ALIGN_PARAGRAPH.LEFT,
                line=1)


def horizontal_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 0.25
    pp = p._p.get_or_add_pPr()
    border = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    for k, v in [('val','single'),('sz','7'),('space','1'),('color','333333')]:
        bottom.set(qn('w:'+k), v)
    border.append(bottom)
    pp.append(border)


def page_heading(doc, page, index):
    day = page['day']
    p = para(doc, f'EXPERIMENT {day:02d}   /   {DATES[day-1]}', 10, True, 3,
             align=WD_ALIGN_PARAGRAPH.LEFT, line=1)
    p.paragraph_format.keep_with_next = True
    p = para(doc, TITLES[day-1], 14, True, 4, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.12)
    p.paragraph_format.keep_with_next = True
    horizontal_rule(doc)
    if index == STARTS[day] - 1:
        p = para(doc, 'Aim: ' + AIMS[day-1], 12, after=6, line=1.5)
        p.paragraph_format.keep_with_next = True
        cautions = {
            2: 'Run each TCP server/client pair separately: all three server variants use port 5000. The calculator variant needs a JDK with JShell.',
            5: 'The supplied program performs in-place all-pairs cost relaxation (Floyd–Warshall style); it does not exchange vectors between routers.',
            7: 'This is a local SMTP-style teaching simulation on TCP port 2525, not a connection to a production mail server.',
            8: 'These sources demonstrate a custom file transfer over TCP port 5000; they do not implement the standard FTP protocol. The repository’s test.txt and received_test.txt both contain “Message transmitted through FTP!”.',
        }
        if day in cautions:
            para(doc, 'Note: ' + cautions[day], 10, after=5, line=1.15)


def add_code(doc, page):
    filename = Path(page['file']).name
    if page['day'] == 1 and page['part'] == 1:
        para(doc, 'This script was written for revised Day 1(b); no shell-script source was supplied in the original repository.',
             10, after=6, line=1.15)
    continuation = f"   /   PART {page['part']} OF {page['total']}" if page['total'] > 1 else ''
    label(doc, f'PROGRAM   /   {filename}{continuation}', 2)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.13)
    p.paragraph_format.right_indent = Inches(0.03)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = font(p.add_run(page['text']), 8.0, face='DejaVu Sans Mono')
    run.font.color.rgb = BLACK
    pp = p._p.get_or_add_pPr()
    shade = OxmlElement('w:shd')
    shade.set(qn('w:fill'), 'F5F5F5')
    pp.append(shade)
    # Docx preserves newlines as w:br and the exact source characters in each chunk.


def terminal_image(lines):
    textfont = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf', 19)
    titlefont = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 16)
    width = max(1050, 55 + max(len(line) for line in lines) * 12)
    height = 65 + len(lines) * 29 + 24
    im = Image.new('RGB', (width, height), '#ececec')
    draw = ImageDraw.Draw(im)
    draw.rectangle((0,0,width,42), fill='#303030')
    draw.text((20,11), 'Terminal transcript • Linux lab', font=titlefont, fill='white')
    for i, line in enumerate(lines):
        draw.text((24, 60 + i * 29), line, font=textfont, fill='#171717')
    stream = BytesIO()
    im.save(stream, 'PNG')
    stream.seek(0)
    return stream


def photo_stream(path):
    # These are crops of the repository's original screenshots, not regenerated outputs.
    im = Image.open(ROOT / path).convert('RGB')
    name = Path(path).name
    if im.width >= 750 and im.height >= 500:
        bottom = 360 if 'SMTPServer' in name else 290
        if name.startswith('LeakyBucket'): bottom = im.height
        im = im.crop((0, 0, im.width, min(bottom, im.height)))
    if path == 'DVR/day6.png':
        im = im.crop((0, 0, im.width, 822))
    stream = BytesIO()
    im.save(stream, 'PNG', optimize=True)
    stream.seek(0)
    return stream, im.size


def picture(doc, stream, size, max_w, max_h):
    w, h = size
    scale = min(max_w / w, max_h / h)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1
    p.add_run().add_picture(stream, width=Inches(w * scale), height=Inches(h * scale))


def images_page(doc, page):
    label(doc, 'SCREENSHOT OF OUTPUT', 3)
    items = page['items']
    for path, caption in items:
        para(doc, caption, 11, True, after=3, align=WD_ALIGN_PARAGRAPH.LEFT, line=1)
        image, size = photo_stream(path)
        picture(doc, image, size, 5.9, 3.8 if len(items) > 1 else 7.9)
    para(doc, 'Evidence: cropped views of the original PNG terminal captures stored in this repository.',
         9, after=0, align=WD_ALIGN_PARAGRAPH.LEFT, line=1, color=GRAY)


def linux_page(doc):
    label(doc, 'TITLE   /   Basics of network configuration files and networking commands in Linux', 2)
    para(doc, 'Inspect /etc/hosts (local name mappings), /etc/resolv.conf (DNS resolvers), and the network manager’s distribution-specific configuration. Use the commands below to inspect addresses, routes, listening sockets and reachability.',
         12, after=7, line=1.5)
    label(doc, 'PROGRAM   /   Linux commands', 4)
    for command, description in [
        ('ip addr show', 'list interfaces and IP addresses'),
        ('ip route show', 'show the routing table and default gateway'),
        ('ping -c 1 127.0.0.1', 'test local connectivity'),
        ('ss -tuln', 'show listening TCP/UDP sockets'),
        ('cat /etc/hosts', 'read local hostname mappings'),
        ('cat /etc/resolv.conf', 'read DNS resolver settings'),
    ]:
        para(doc, command + '  —  ' + description, 10.5, after=1, line=1.1,
             align=WD_ALIGN_PARAGRAPH.LEFT)
    label(doc, 'SCREENSHOT OF OUTPUT', 8)
    out = subprocess.run(['bash', '-c', "ip -br addr show lo; ip route get 127.0.0.1; ping -c 1 -W 2 127.0.0.1"],
                         text=True, capture_output=True, check=True).stdout.splitlines()
    lines = ['$ ip -br addr show lo', out[0], '$ ip route get 127.0.0.1'] + out[1:3]
    lines += ['$ ping -c 1 -W 2 127.0.0.1'] + out[3:]
    picture(doc, terminal_image(lines), (1050, 65 + len(lines)*29+24), 5.8, 3.5)
    para(doc, 'Rendered transcript of commands executed on the local loopback interface; not a claim of campus-network configuration.',
         9, after=0, align=WD_ALIGN_PARAGRAPH.LEFT, line=1, color=GRAY)


def dryrun_page(doc):
    label(doc, 'SCREENSHOT OF OUTPUT   /   network_setup.sh --dry-run', 3)
    args = ['bash', str(ROOT / 'Day 1/network_setup.sh'), '--dry-run', 'eth0',
            '192.0.2.10/24', '192.0.2.1', '192.0.2.1']
    text = subprocess.run(args, text=True, capture_output=True, check=True).stdout.splitlines()
    lines = ['$ bash network_setup.sh --dry-run eth0 192.0.2.10/24 192.0.2.1 192.0.2.1'] + text
    picture(doc, terminal_image(lines), (max(1050,55+max(map(len,lines))*12), 65+len(lines)*29+24), 5.9, 4.4)
    para(doc, 'Dry-run evidence only: the script prints the three configuration commands but does not change an interface, run ping or write a route file. Run with sudo on a lab machine to perform and capture those steps.',
         12, after=9, line=1.5)
    label(doc, 'RUN ON A LAB MACHINE', 4)
    para(doc, 'sudo bash network_setup.sh eth0 192.0.2.10/24 192.0.2.1 192.0.2.1 routes.txt',
         10, after=8, align=WD_ALIGN_PARAGRAPH.LEFT, line=1)
    para(doc, 'Use addresses and an interface supplied by your instructor. Changing the default route can disconnect a remote session.',
         10, after=0, line=1.2)


FAMILIAR = {
    9: {
        'overview': 'Wireshark is a free, open-source network protocol analyzer used to capture and examine packets as data travels through a network.',
        'features': [
            'Packet capture: records packets transmitted and received by a computer.',
            'Protocol analysis: examines TCP, UDP, HTTP, DNS and ICMP traffic.',
            'Packet details: shows source and destination IP addresses, ports, protocol and packet length.',
            'Display filters: isolates traffic of interest, such as tcp, udp or icmp.',
            'Troubleshooting: helps identify loss, retransmissions and connection failures.',
        ],
        'activity': 'Select a network interface, start a capture, apply a display filter and inspect a packet’s decoded fields.',
        'result': 'The Wireshark familiarization material describes capturing and analyzing packets with suitable filters to understand packet details and network protocols.',
    },
    10: {
        'overview': 'NS2 (Network Simulator 2) is a discrete-event simulator for designing and analyzing computer networks without a physical network.',
        'features': [
            'Network simulation: models virtual nodes, links and connections.',
            'Protocol study: investigates TCP, UDP, routing and MAC protocols.',
            'Tcl scripts: define scenarios and schedule simulated events.',
            'NAM: animates packet transmission and other network activity.',
            'Trace files: record transmission, reception and loss; support analysis of throughput, delay, loss and packet delivery ratio.',
        ],
        'activity': 'Study a simple Tcl network scenario and inspect its trace file or NAM animation.',
        'result': 'The NS2 familiarization material describes a basic simulation and the use of topologies, packet traces and performance analysis.',
    },
    11: {
        'overview': 'Cisco Packet Tracer is a simulation tool for building virtual networks without physical equipment. It provides routers, switches, PCs and servers that can be interconnected and configured.',
        'features': [
            'Network topology: place and connect virtual devices.',
            'Addressing and configuration: assign IP addresses and configure devices.',
            'Connectivity testing: use ping from a PC command prompt.',
            'Packet observation: watch data move through the virtual network.',
        ],
        'activity': 'Build a small two-PC network, assign addresses in one subnet and test connectivity with ping.',
        'result': 'The Packet Tracer familiarization material describes building a basic network, configuring devices and testing connectivity to understand addressing, topology and packet transmission.',
    },
}


def familiarization_page(doc, page):
    n = page['day']
    material = FAMILIAR[n]
    label(doc, 'TITLE   /   ' + TITLES[n-1], 2)
    label(doc, 'STUDY MATERIAL   /   Familiarisation.docx', 10)
    para(doc, material['overview'], 12, after=7, line=1.5)
    label(doc, 'KEY FEATURES', 7)
    for item in material['features']:
        para(doc, '•  ' + item, 12, after=4, line=1.5)
    label(doc, 'FAMILIARIZATION ACTIVITY', 8)
    para(doc, material['activity'], 12, after=5, line=1.5)
    label(doc, 'RESULT / LEARNING OUTCOME', 8)
    para(doc, material['result'], 12, after=0, line=1.5)


def set_num_start(section, value):
    sect_pr = section._sectPr
    pg = OxmlElement('w:pgNumType')
    pg.set(qn('w:start'), str(value))
    cols = sect_pr.find(qn('w:cols'))
    sect_pr.insert(list(sect_pr).index(cols), pg) if cols is not None else sect_pr.append(pg)


def add_field(paragraph, instruction):
    r = paragraph.add_run()
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin'); r._r.append(b)
    r = paragraph.add_run()
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = ' ' + instruction + ' '
    r._r.append(it)
    r = paragraph.add_run()
    s = OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'), 'separate'); r._r.append(s)
    font(paragraph.add_run('1'), 9)
    r = paragraph.add_run()
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r._r.append(e)


def make_cover(doc):
    with ZipFile(ROOT / 'Computer Networks Record fmt.docx') as z:
        logo = BytesIO(z.read('word/media/image1.png'))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.add_run().add_picture(logo, height=Inches(1.45))
    para(doc, '102003/CS500A  •  COMPUTER NETWORKS', 11, True, before=18, after=11,
         align=WD_ALIGN_PARAGRAPH.CENTER, line=1)
    para(doc, 'COMPUTER NETWORKS\nLAB CODEBOOK', 20, True, after=17,
         align=WD_ALIGN_PARAGRAPH.CENTER, line=1.1)
    horizontal_rule(doc)
    para(doc, 'Submitted in partial fulfillment of the requirements for the award of the degree of',
         12, after=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, 'Bachelor of Technology\nin\nComputer Science and Engineering',
         14, True, after=26, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.6)
    para(doc, 'Name:  _______________________________________', 12, after=7,
         align=WD_ALIGN_PARAGRAPH.LEFT)
    para(doc, 'Roll No.:  ______________    University Reg. No.:  ___________________',
         12, after=23, align=WD_ALIGN_PARAGRAPH.LEFT)
    para(doc, 'Department of Computer Science and Engineering\nRajagiri School of Engineering & Technology (Autonomous)\n(Parent University: APJ Abdul Kalam Technological University)\nRajagiri Valley, Kakkanad, Kochi – 682039',
         11, after=16, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.35)
    para(doc, 'September 2026', 12, True, align=WD_ALIGN_PARAGRAPH.CENTER)


def make_index(doc):
    para(doc, 'INDEX  /  REVISED LAB CYCLE', 14, True, before=10, after=5,
         align=WD_ALIGN_PARAGRAPH.LEFT, line=1)
    horizontal_rule(doc)
    para(doc, 'Experiment pages begin at 1; cover and index are unnumbered. Dates follow the supplied reference record.',
         10, after=10, line=1.2)
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    table.autofit = False
    for c, width in zip(table.columns, [0.55, 4.30, 1.05, 0.8]):
        c.width = Inches(width)
    headings = ['Sl. No.', 'Experiment name (revised lab cycle)', 'Date', 'Page No.']
    for cell, text in zip(table.rows[0].cells, headings):
        cell.text = text
        cell._tc.get_or_add_tcPr().append(OxmlElement('w:shd'))
        cell._tc.tcPr[-1].set(qn('w:fill'), 'E8E8E8')
    index_names = [
        'Familiarization of (a) Linux network configuration files and commands; (b) shell scripting for IP, gateway, connectivity and routing',
        'Implement client–server communication using socket programming and TCP',
        'Implement client–server communication using socket programming and UDP',
        'Implement Stop-and-Wait, Go-Back-N and Selective Repeat ARQ',
        'Implement and simulate Distance Vector Routing',
        'Implement and simulate the Leaky Bucket algorithm',
        'Implement Simple Mail Transfer Protocol',
        'Implement File Transfer Protocol',
        'Introduction to Wireshark',
        'Familiarization with NS2',
        'Study of Packet Tracer',
    ]
    for n in range(1,12):
        row = table.add_row().cells
        for cell, text in zip(row, [str(n), index_names[n-1], DATES[n-1], str(STARTS[n])]):
            cell.text = text
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(5)
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.line_spacing = 1.12
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 1 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    font(run, 9.5, ri == 0)
    para(doc, 'Experiments 9–11 are familiarization records based on Familiarisation.docx; no screenshots are required for those three entries.',
         9.5, before=10, after=0, line=1.15, color=GRAY)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin, sec.bottom_margin = Inches(.78), Inches(.74)
    sec.left_margin, sec.right_margin = Inches(.76), Inches(.76)
    sec.header_distance, sec.footer_distance = Inches(.36), Inches(.35)
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'; st.font.size = Pt(12); st.font.color.rgb = BLACK
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(5)
    make_cover(doc)
    doc.add_page_break()
    make_index(doc)
    exp = doc.add_section(WD_SECTION_START.NEW_PAGE)
    exp.page_width, exp.page_height = sec.page_width, sec.page_height
    exp.top_margin, exp.bottom_margin = sec.top_margin, sec.bottom_margin
    exp.left_margin, exp.right_margin = sec.left_margin, sec.right_margin
    exp.header_distance, exp.footer_distance = sec.header_distance, sec.footer_distance
    exp.header.is_linked_to_previous = False
    exp.footer.is_linked_to_previous = False
    set_num_start(exp, 1)
    header = exp.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(header.add_run('COMPUTER NETWORKS  /  LAB CODEBOOK'), 8.5, color=GRAY)
    footer = exp.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run('CN  •  '), 9, color=GRAY)
    add_field(footer, 'PAGE')
    for i, page in enumerate(PAGES):
        if i: doc.add_page_break()
        page_heading(doc, page, i)
        kind = page['kind']
        if kind == 'code': add_code(doc, page)
        elif kind == 'images': images_page(doc, page)
        elif kind == 'linux': linux_page(doc)
        elif kind == 'dryrun': dryrun_page(doc)
        else: familiarization_page(doc, page)
    # Ask Word to refresh page fields when opening the file.
    settings = doc.settings.element
    update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true')
    settings.append(update)
    doc.core_properties.title = 'Computer Networks Lab Codebook — Revised Lab Cycle'
    doc.core_properties.subject = 'CN Lab Record, 11 experiments, source code and output evidence'
    doc.core_properties.keywords = 'Computer Networks; Lab; TCP; UDP; ARQ; routing; SMTP; FTP'
    doc.save(DEST)
    print('Created', DEST, f'({len(PAGES)} numbered experiment pages)')
    print('Index starts:', STARTS)


if __name__ == '__main__':
    main()
